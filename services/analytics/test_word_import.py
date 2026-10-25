import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
import main
import word_import


def document_bytes(rows, extra=None):
    document = Document()
    document.add_paragraph('Описание источника: не является записью')
    for values in [rows] + ([extra] if extra else []):
        table = document.add_table(rows=0, cols=len(values[0]))
        for row in values:
            for cell, value in zip(table.add_row().cells, row): cell.text = str(value)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class WordImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.override = patch.object(main, 'DB', Path(self.temp.name) / 'test.sqlite3')
        self.override.start()
        main.initialize()
        self.client = TestClient(main.app)

    def tearDown(self):
        self.client.close()
        self.override.stop()
        self.temp.cleanup()

    def upload(self, kind, raw, name='data.docx'):
        return self.client.post('/import/' + kind, params={'filename': name}, content=raw)

    def sample(self, kind):
        row = main.records(kind)[0].copy()
        for key in ('id', 'demo', 'updatedAt', 'distance'): row.pop(key, None)
        row['name'] = 'Word import test ' + kind
        row['source'] = 'Таблица Word'
        row['sourceUrl'] = 'https://example.org/word-import'
        return row

    def test_import_all_kinds_and_repeat_without_duplicates(self):
        for kind in ('companies', 'projects', 'deals'):
            with self.subTest(kind=kind):
                sample = self.sample(kind)
                raw = document_bytes([list(sample), list(sample.values())])
                before = len(main.records(kind))
                result = self.upload(kind, raw)
                self.assertEqual(result.status_code, 200, result.text)
                self.assertEqual(result.json()['imported'], 1)
                self.assertEqual(len(main.records(kind)), before + 1)
                self.assertEqual(self.upload(kind, raw).status_code, 200)
                self.assertEqual(len(main.records(kind)), before + 1)

    def test_invalid_later_row_is_atomic(self):
        sample = self.sample('companies')
        invalid = dict(sample, name='Invalid', inn='123')
        before = len(main.records('companies'))
        result = self.upload('companies', document_bytes([list(sample), list(sample.values()), list(invalid.values())]))
        self.assertEqual(result.status_code, 422)
        self.assertEqual(len(main.records('companies')), before)

    def test_reordered_columns_in_second_table(self):
        sample = self.sample('companies')
        second = dict(sample, name='Second Word record', inn='0265052371')
        keys = list(reversed(second))
        raw = document_bytes([list(sample), list(sample.values())], [keys, [second[key] for key in keys]])
        result = self.upload('companies', raw)
        self.assertEqual(result.status_code, 200, result.text)
        imported = [r for r in main.records('companies') if r.get('inn') == '0265052371']
        self.assertEqual(imported[0]['inn'], '0265052371')
        self.assertEqual(result.json()['imported'], 2)

    def test_bad_file_plain_text_and_empty_tables(self):
        for raw in (b'not a docx', document_bytes([['name', 'region']])):
            self.assertEqual(self.upload('companies', raw).status_code, 422)
        document = Document()
        document.add_paragraph('Company name without table')
        output = io.BytesIO(); document.save(output)
        self.assertEqual(self.upload('companies', output.getvalue()).status_code, 422)

    def test_merged_cells_rejected(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = 'name'
        output = io.BytesIO(); document.save(output)
        self.assertEqual(self.upload('companies', output.getvalue()).status_code, 422)

    def test_word_template_can_be_filled_and_imported(self):
        for kind in ('companies', 'projects', 'deals'):
            response = self.client.get('/template/' + kind + '?format=docx')
            self.assertEqual(response.status_code, 200)
            document = Document(io.BytesIO(response.content))
            table = document.tables[0]
            sample = self.sample(kind)
            for cell, header in zip(table.add_row().cells, table.rows[0].cells):
                cell.text = str(sample[header.text])
            output = io.BytesIO(); document.save(output)
            result = self.upload(kind, output.getvalue())
            self.assertEqual(result.status_code, 200, result.text)

    def test_doc_uses_converter_and_reports_unavailable(self):
        with patch.object(word_import.subprocess, 'run', side_effect=FileNotFoundError):
            result = self.upload('companies', b'legacy document', 'data.doc')
            self.assertEqual(result.status_code, 422)
            self.assertIn('DOCX', result.json()['detail'])

    def test_row_limit(self):
        raw = document_bytes([['name'], ['one'], ['two']])
        with self.assertRaises(ValueError): word_import.word_to_csv(raw, 'data.docx', max_rows=1)


if __name__ == '__main__':
    unittest.main()
