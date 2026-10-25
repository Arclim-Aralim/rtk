"""Read Word tables without executing document code or guessing missing fields."""
import csv
import io
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from docx import Document


def read_word_tables(raw, filename, max_rows=10_000):
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.doc', '.docx'):
        raise ValueError('Поддерживаются документы DOC и DOCX')
    if suffix == '.doc':
        # Separate temporary profile allows concurrent conversions without a shared lock.
        with tempfile.TemporaryDirectory(prefix='word-import-') as folder:
            root = Path(folder)
            source = root / 'upload.doc'
            source.write_bytes(raw)
            try:
                result = subprocess.run(
                    ['libreoffice', '-env:UserInstallation=' + (root / 'profile').as_uri(),
                     '--headless', '--convert-to', 'docx', '--outdir', folder, str(source)],
                    capture_output=True, timeout=45, check=False,
                )
            except FileNotFoundError as exc:
                raise ValueError('Конвертер DOC недоступен. Сохраните документ как DOCX') from exc
            except subprocess.TimeoutExpired as exc:
                raise ValueError('Превышено время чтения DOC. Сохраните документ как DOCX') from exc
            target = root / 'upload.docx'
            if result.returncode or not target.exists():
                raise ValueError('Не удалось прочитать DOC. Проверьте файл и отсутствие пароля')
            if target.stat().st_size > 25_000_000:
                raise ValueError('Слишком большой документ после преобразования DOC')
            raw = target.read_bytes()
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 50_000_000:
                raise ValueError('Слишком большой распакованный документ Word')
            if 'word/document.xml' not in archive.namelist():
                raise ValueError('Файл не является документом Word')
        document = Document(io.BytesIO(raw))
        if document.element.xpath('//w:vMerge') or document.element.xpath('//w:gridSpan'):
            raise ValueError('Разделите объединённые ячейки в таблице Word')
        tables, count = [], 0
        for index, table in enumerate(document.tables, 1):
            rows = []
            for row in table.rows:
                cells = row.cells
                if any(cell.tables for cell in cells):
                    raise ValueError('Вложенные таблицы Word не поддерживаются; перенесите их на верхний уровень')
                if len(cells) > 100:
                    raise ValueError('В таблице Word допускается не более 100 столбцов')
                if len({id(cell._tc) for cell in cells}) != len(cells):
                    raise ValueError('Разделите объединённые ячейки в таблице Word')
                values = [cell.text.strip() for cell in cells]
                if not any(values):
                    continue
                count += 1
                if count > max_rows + len(document.tables):
                    raise ValueError(f'Допускается не более {max_rows} строк данных')
                if any(len(value) > 10_000 for value in values):
                    raise ValueError('Слишком длинный текст в ячейке Word')
                rows.append(values)
            if rows:
                tables.append((f'Таблица {index}', rows))
        if not tables:
            raise ValueError('В документе нет таблиц. Используйте шаблон Word; текст и сканы не импортируются')
        return tables
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError('Повреждённый или защищённый документ Word') from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Не удалось прочитать таблицы Word. Проверьте структуру документа') from exc


def word_to_csv(raw, filename, max_rows=10_000):
    tables = read_word_tables(raw, filename, max_rows)
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    header, total = None, 0
    for name, rows in tables:
        current = rows[0]
        if any(not key for key in current) or len(set(current)) != len(current):
            raise ValueError(f'{name}: первая строка должна содержать уникальные названия полей из шаблона')
        if header is None:
            header = current
            writer.writerow(header)
        elif set(current) != set(header):
            raise ValueError(f'{name}: все таблицы должны содержать одинаковые поля выбранного типа данных')
        for row in rows[1:]:
            if row == current:
                continue
            if len(row) != len(current):
                raise ValueError(f'{name}: число ячеек не совпадает с заголовком')
            values = dict(zip(current, row))
            writer.writerow([values[key] for key in header])
            total += 1
            if total > max_rows:
                raise ValueError(f'Допускается не более {max_rows} строк данных')
    return output.getvalue().encode('utf-8')
