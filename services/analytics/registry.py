"""Source adapters: preserve unknowns, exact INN joins, registration vs operating sites."""
import csv
import hashlib
import io
import math
import re
from datetime import date, datetime
from zipfile import ZipFile, BadZipFile

FNS_URL = 'https://rmsp.nalog.ru/search.html?mode=extended'
OTHER_INDUSTRY = 'Прочие отрасли'


def text(value):
    return '' if value is None else str(value).strip()


def identifier(value):
    if isinstance(value, float) and value.is_integer(): value = int(value)
    result = text(value)
    if not re.fullmatch(r'(?:[0-9]{10}|[0-9]{12})', result):
        raise ValueError('ИНН должен содержать 10 или 12 цифр; ведущие нули необходимо сохранить')
    return result


def company_id(inn):
    return 'real-' + hashlib.sha256(inn.encode()).hexdigest()[:20]


def iso_date(value):
    if not value: return None
    if isinstance(value, (date, datetime)): return value.strftime('%Y-%m-%d')
    value = text(value)
    for pattern in ('%d.%m.%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
        try: return datetime.strptime(value, pattern).date().isoformat()
        except ValueError: pass
    raise ValueError(f'неверная дата: {value}')


def region_name(value):
    return re.sub(r'^\d{1,2}\s*[-—]?\s*', '', text(value))


def industry_from_okved(value):
    match = re.match(r'\s*(\d{2})', text(value))
    code = int(match[1]) if match else -1
    if 1 <= code <= 3: return 'АПК'
    if 5 <= code <= 9 or code == 35: return 'ТЭК'
    if 10 <= code <= 33: return 'Промышленность'
    if 41 <= code <= 43: return 'Строительство'
    if 45 <= code <= 47: return 'Торговля'
    if 49 <= code <= 53: return 'Логистика'
    if 55 <= code <= 56 or 85 <= code <= 96: return 'Услуги населению'
    return OTHER_INDUSTRY


def finite(value):
    if value in (None, ''): return None
    number = float(text(value).replace(' ', '').replace('\xa0', '').replace(',', '.'))
    if not math.isfinite(number): raise ValueError('число должно быть конечным')
    return number


def read_tables(raw, filename='upload.csv', max_rows=50_000):
    """CSV or XLSX. Ignore preamble rows; bound ZIP expansion and row/column counts."""
    if raw[:2] == b'PK':
        from openpyxl import load_workbook
        try:
            with ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 1_500_000_000:
                    raise ValueError('слишком большой распакованный XLSX')
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            result = []
            total = 0
            for sheet in workbook:
                if (sheet.max_column or 0) > 100: raise ValueError('слишком много столбцов')
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if any(v is not None for v in row): rows.append(row)
                    total += 1
                    if total > max_rows: raise ValueError(f'не более {max_rows} строк за импорт')
                result.append((sheet.title, rows))
            workbook.close()
            return result
        except (BadZipFile, KeyError, OSError) as exc:
            raise ValueError('повреждённый или неподдерживаемый XLSX') from exc
    try: content = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        try: content = raw.decode('cp1251')
        except UnicodeDecodeError as exc: raise ValueError('CSV должен быть UTF-8 или Windows-1251') from exc
    lines = content.splitlines()
    sample = '\n'.join(lines[:10])
    try: dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
    except csv.Error: dialect = csv.excel
    rows = list(csv.reader(io.StringIO(content), dialect))
    if len(rows) > max_rows: raise ValueError(f'не более {max_rows} строк за импорт')
    return [(filename, rows)]


def dict_rows(rows):
    # FNS exports can contain a title before the column headings.
    for offset, row in enumerate(rows[:20]):
        names = [text(v) for v in row]
        if 'ИНН' in names or 'inn' in names:
            return [(offset+i+2, dict(zip(names, values))) for i, values in enumerate(rows[offset+1:]) if any(text(v) for v in values)]
    raise ValueError('не найден столбец ИНН / inn')


def parse_companies(raw, filename, is_client=False):
    parsed, seen = [], set()
    tables = read_tables(raw, filename)
    suitable = [(name, rows) for name, rows in tables if any('ИНН' in [text(v) for v in row] or 'inn' in [text(v) for v in row] for row in rows[:20])]
    if not suitable: raise ValueError('файл не содержит таблицы с ИНН')
    for sheet, rows in suitable:
        for line, row in dict_rows(rows):
            try:
                inn = identifier(row.get('ИНН', row.get('inn')))
                if inn in seen: raise ValueError(f'повтор ИНН {inn}; выберите один актуальный срез реестра')
                seen.add(inn)
                name = text(row.get('Наименование / ФИО', row.get('Наименование субъекта', row.get('name'))))
                if not name: name = 'ИНН ' + inn
                region = region_name(row.get('Регион', row.get('region'))) or None
                city = text(row.get('Город', row.get('registrationLocality')))
                locality = text(row.get('Населенный пункт'))
                district = text(row.get('Район'))
                registered = ', '.join(v for v in (region, district, city, locality) if v)
                excluded = iso_date(row.get('Дата исключения из реестра', row.get('mspExcludedAt')))
                included = iso_date(row.get('Дата включения в реестр', row.get('mspIncludedAt')))
                termination = iso_date(row.get('terminationDate'))
                activity = text(row.get('activityStatus')) or ('terminated' if termination else 'unknown')
                if activity not in ('active', 'terminated', 'unknown'): raise ValueError('activityStatus: active / terminated / unknown')
                okved = text(row.get('Основной вид деятельности', row.get('okved')))
                employees = finite(row.get('Среднесписочная численность работников за предшествующий календарный год', row.get('employees')))
                lat, lon = finite(row.get('lat')), finite(row.get('lon'))
                if (lat is None) != (lon is None) or (lat is not None and not (-90 <= lat <= 90 and -180 <= lon <= 180)):
                    raise ValueError('некорректные координаты')
                parsed.append(dict(id=company_id(inn), inn=inn, name=name, region=region, registrationRegion=region, registrationLocality=locality or city or None, registrationDistrict=district or None, registrationAddress=text(row.get('registrationAddress')) or registered or None, address=registered or None, industry=industry_from_okved(okved), okved=okved, ogrn=text(row.get('ОГРН', row.get('ogrn'))), entityType=text(row.get('Тип субъекта')), category=text(row.get('Категория')), employees=employees, revenue=finite(row.get('revenue')), growth=finite(row.get('growth')), client=True if is_client else None, clientMatch='confirmed' if is_client else 'unknown', clientSource=filename if is_client else None, mspStatus='excluded' if excluded else 'present' if included else 'unknown', mspIncludedAt=included, mspExcludedAt=excluded, activityStatus=activity, terminationDate=termination, statusSource=text(row.get('statusSource')) or None, statusCheckedAt=iso_date(row.get('statusCheckedAt')), lat=lat, lon=lon, geoAccuracy='provided' if lat is not None else 'unknown', source=f'{filename} · {sheet}' if is_client else 'ФНС · Реестр МСП', sourceUrl='' if is_client else FNS_URL, updatedAt=date.today().isoformat(), demo=False, sourceRow=line))
            except (ValueError, TypeError) as exc:
                raise ValueError(f'{sheet}, строка {line}: {exc}') from exc
    if not parsed: raise ValueError('нет записей для импорта')
    return parsed


def parse_sites(raw, filename):
    result, seen = [], set()
    for sheet, rows in read_tables(raw, filename, max_rows=10_000):
        for line, row in dict_rows(rows):
            try:
                inn = identifier(row.get('inn', row.get('ИНН')))
                region = region_name(row.get('region'))
                address, source = text(row.get('address')), text(row.get('source'))
                locality = text(row.get('locality'))
                status = text(row.get('status')) or 'candidate'
                url = text(row.get('sourceUrl'))
                if not region or not address or not locality or not source or not url.startswith(('http://', 'https://')):
                    raise ValueError('обязательны region, locality, address, source и HTTP(S) sourceUrl')
                if status not in ('confirmed', 'candidate', 'rejected'): raise ValueError('status: confirmed / candidate / rejected')
                observed = iso_date(row.get('observedAt'))
                if not observed: raise ValueError('укажите observedAt — дату подтверждения источника')
                lat, lon = finite(row.get('lat')), finite(row.get('lon'))
                if (lat is None) != (lon is None) or (lat is not None and not (-90 <= lat <= 90 and -180 <= lon <= 180)):
                    raise ValueError('некорректные координаты')
                key = hashlib.sha256((inn+'|'+address.casefold()+'|'+url).encode()).hexdigest()[:24]
                if key in seen: raise ValueError('дубликат площадки')
                seen.add(key)
                result.append(dict(id='site-'+key, inn=inn, region=region, locality=locality, address=address, source=source, sourceUrl=url, observedAt=observed, status=status, lat=lat, lon=lon, geoAccuracy=text(row.get('geoAccuracy')) or 'provided', matchMethod='exact_inn', demo=False))
            except (ValueError, TypeError) as exc: raise ValueError(f'{sheet}, строка {line}: {exc}') from exc
    return result


def normalize_locality(value):
    value = text(value).lower().replace('ё', 'е')
    value = re.sub(r'^(?:г\.|город|п\.|поселок|с\.|село|ст-ца)\s*', '', value)
    return re.sub(r'[^а-яa-z0-9]+', '', value)


def enrich_company(company, sites, mode='registration', region='', confidence='confirmed'):
    result = company.copy()
    candidates = sites.get(company.get('inn'), [])
    confirmed = [s for s in candidates if s['status'] == 'confirmed']
    usable = [s for s in candidates if s['status'] != 'rejected' and (confidence == 'all' or s['status'] == confidence)]
    result['operatingSites'] = candidates
    registered = company.get('registrationRegion') or company.get('region')
    locality = normalize_locality(company.get('registrationLocality'))
    result['crossRegion'] = bool(registered and any(s['region'] != registered for s in confirmed))
    result['crossLocality'] = bool(locality and any(normalize_locality(s.get('locality')) and (s['region'] != registered or normalize_locality(s['locality']) != locality) for s in confirmed))
    result['siteCandidates'] = sum(s['status'] == 'candidate' for s in candidates)
    registration_matches = not region or registered == region
    operating_matches = any(not region or s['region'] == region for s in usable)
    result['_regionMatches'] = registration_matches if mode == 'registration' else operating_matches if mode == 'operations' else registration_matches or operating_matches
    locations = []
    if mode in ('registration', 'both') and company.get('lat') is not None and company.get('lon') is not None and registration_matches:
        locations.append(dict(id='registration', lat=company['lat'], lon=company['lon'], label=company.get('registrationAddress') or company.get('address'), accuracy=company.get('geoAccuracy', 'provided'), kind='registration'))
    if mode in ('operations', 'both'):
        for s in usable:
            if (not region or s['region'] == region) and s.get('lat') is not None and s.get('lon') is not None:
                locations.append(dict(id=s['id'], lat=s['lat'], lon=s['lon'], label=s['address'], accuracy=s.get('geoAccuracy', 'provided'), kind='operations', status=s['status']))
    result['mapLocations'] = locations
    return result


def parse_facts(raw, filename):
    tables = read_tables(raw, filename)
    companies = parse_companies(raw, filename, is_client=True)
    projects, items, regional, levels, works = [], [], [], [], []
    for sheet, rows in tables:
        if not rows: continue
        headers = [text(v) for v in rows[0]]
        if 'Инициатор инвестиционного проекта' in headers:
            for line, values in enumerate(rows[1:], 2):
                r = dict(zip(headers, values))
                name = text(r.get('Название и суть инвестиционного проекта (в том числе и мощность)'))
                if not name: continue
                initiator = text(r.get('Инициатор инвестиционного проекта'))
                municipality = text(r.get('Муниципальное образование'))
                end = text(r.get('Ввод в эксплуатацию'))
                # Source contains only years; do not invent an exact event date or actual construction stage.
                if end.endswith('.0'): end = end[:-2]
                key = hashlib.sha256((name+initiator+municipality).encode()).hexdigest()[:24]
                value = (name+' '+text(r.get('Куратор проекта'))).lower()
                industry = 'АПК' if any(v in value for v in ['агро','комбикорм','сельск','рыб','рак','птиц','молоч']) else 'Логистика' if any(v in value for v in ['логист','склад','транспорт']) else 'Промышленность' if any(v in value for v in ['завод','производств','промышлен']) else 'Услуги населению'
                projects.append(dict(id='project-'+key, name=name, initiator=initiator, region='Ростовская область', municipality=municipality, industry=industry, stage='Не уточнена', need='ИТ-инфраструктура', needStatus='suggested', investment=finite(r.get('Всего (млн руб.)')), date=end if re.fullmatch(r'\d{4}',end) else '', datePrecision='year' if re.fullmatch(r'\d{4}',end) else 'unknown', startYear=text(r.get('Начало')), newJobs=finite(r.get('Количество новых рабочих мест (мест)')), curator=text(r.get('Куратор проекта')), lat=None, lon=None, source=f'{filename} · {sheet}', sourceUrl='', updatedAt=date.today().isoformat(), demo=False))
        elif 'Компонент (прибор/узел)' in headers:
            for values in rows[1:]:
                r = dict(zip(headers,values))
                if not r.get('ID'): continue
                items.append(dict(id=text(r['ID']),group=text(r['Группа системы']),name=text(r['Компонент (прибор/узел)']),unit=text(r['Ед. изм.']),price=finite(r['Базовая цена (руб.)']),complexity=finite(r['Коэф. сложности'])))
        elif 'Регион (примеры)' in headers:
            regional = [dict(name=text(v[0]),factor=finite(v[1]),note=text(v[2])) for v in rows[1:] if v[0]]
        elif 'Уровень проекта' in headers:
            levels = [dict(name=text(v[0]),factor=finite(v[1]),note=text(v[2])) for v in rows[1:] if v[0]]
        elif 'Этап работ' in headers:
            works = [dict(name=text(v[0]),rate=float(text(v[1]).replace('%',''))/100,note=text(v[2])) for v in rows[1:] if v[0]]
    if not projects or not items: raise ValueError('в фактуре не найдены инвестиционные проекты или каталог оборудования')
    return dict(companies=companies, projects=projects, pricing=dict(id='provided',items=items,regional=regional,levels=levels,works=works,source=filename,updatedAt=date.today().isoformat(),priceStatus='provided_not_market_verified'))
