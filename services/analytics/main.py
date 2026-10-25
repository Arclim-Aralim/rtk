"""Analytical microservice. Internal API, exposed only through the authenticated gateway."""
import csv
import io
import json
import math
import os
import sqlite3
from contextlib import contextmanager
from datetime import date, timedelta
from pathlib import Path
from typing import Literal
from starlette.concurrency import run_in_threadpool

from fastapi import FastAPI, HTTPException, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel, Field
import registry
import geocoding
import word_import

DATA = Path(os.getenv('DATA_DIR', './data'))
DATA.mkdir(parents=True, exist_ok=True)
DB = DATA / 'analytics.sqlite3'
app = FastAPI(title='Ростелеком · Аналитика', version='1.0.0')
REGIONS = {'Ростовская область': (47.23, 39.72, 1.0), 'Краснодарский край': (45.04, 38.98, 1.12), 'Волгоградская область': (48.71, 44.51, .95), 'Воронежская область': (51.67, 39.2, 1.03)}
INDUSTRIES = ['Промышленность', 'АПК', 'Торговля', 'ТЭК', 'Строительство', 'Логистика', 'Услуги населению', 'Прочие отрасли']
NEEDS = {'Промышленность': 'АСУ ТП', 'АПК': 'Видеонаблюдение', 'Торговля': 'ИТ-инфраструктура', 'ТЭК': 'Кибербезопасность', 'Строительство': 'Связь', 'Логистика': 'ИТ-инфраструктура'}
STAGES = ['Проектирование', 'Строительство', 'Модернизация', 'Запуск', 'Не уточнена']
TYPES = ['ИТ-инфраструктура', 'Видеонаблюдение', 'Кибербезопасность', 'АСУ ТП', 'Связь']


@contextmanager
def connect():
    con = sqlite3.connect(DB, timeout=20)
    con.row_factory = sqlite3.Row
    try:
        with con:
            yield con
    finally:
        con.close()


def initialize():
    with connect() as con:
        con.execute('CREATE TABLE IF NOT EXISTS records (kind TEXT NOT NULL, id TEXT NOT NULL, body TEXT NOT NULL, PRIMARY KEY(kind,id))')
        con.execute('CREATE TABLE IF NOT EXISTS metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)')
        if con.execute("SELECT 1 FROM metadata WHERE key='initialized'").fetchone():
            return
        names = ['Южный промышленный парк', 'Агрокомплекс Дон', 'Торговая сеть Горизонт', 'Энергия региона', 'Девелопмент Юг', 'Логистический центр Вектор', 'Завод новых материалов', 'Агрохолдинг Поле', 'Ритейл Партнёр', 'Тепловые системы', 'Стройинвест', 'Транспортный узел']
        for r, (region, (lat, lon, _)) in enumerate(REGIONS.items()):
            for i, name in enumerate(names):
                n = r * 12 + i + 1
                industry = INDUSTRIES[i % 6]
                company = dict(id=f'c{n}', name=f'ООО «{name} {r+1}»', inn=f'9900{n:06d}', region=region, industry=industry, revenue=round(120 + (n * 167) % 4500, 1), growth=round(-8 + (n * 7.3) % 49, 1), client=i % 4 == 0, lat=round(lat + math.sin(n * 2.3) * .65, 5), lon=round(lon + math.cos(n * 1.7) * .95, 5), address=f'{region}, демонстрационная площадка {i+1}', employees=50+n*13, source='Синтетический набор для демонстрации', sourceUrl='', updatedAt=date.today().isoformat(), demo=True)
                put(con, 'companies', company)
            for i in range(4):
                n = r * 4 + i + 1
                industry = INDUSTRIES[(r+i) % 6]
                project = dict(id=f'p{n}', name=['Строительство производственного комплекса', 'Модернизация распределительного центра', 'Запуск агропромышленной площадки', 'Развитие индустриального парка'][i], initiator=f'ООО «{names[(r+i) % 12]}»', region=region, industry=industry, stage=STAGES[i], need=NEEDS[industry], investment=round(350+n*125.5, 1), date=(date.today()+timedelta(days=14+n*9)).isoformat(), lat=lat+(i-1.5)*.15, lon=lon+(i-1.5)*.23, source='Синтетический набор для демонстрации', sourceUrl='', updatedAt=date.today().isoformat(), demo=True)
                put(con, 'projects', project)
        for i in range(40):
            kind = TYPES[i % 5]
            units = 15 + (i * 13) % 180
            cost = units * [42000, 18000, 65000, 85000, 24000][i % 5]
            put(con, 'deals', dict(id=f'd{i+1}', name=f'{kind} · объект {i+1}', type=kind, region=list(REGIONS)[i % 4], units=units, months=3+i%8, price=round(cost*(1.25+(i%4)*.06)), cost=cost, date=f'2025-{1+i%12:02d}-15', source='Синтетическая история сделок', demo=True))
        con.execute("INSERT INTO metadata VALUES ('initialized', '1')")


def put(con, kind, record):
    con.execute('INSERT OR REPLACE INTO records VALUES (?,?,?)', (kind, record['id'], json.dumps(record, ensure_ascii=False)))


def records(kind):
    with connect() as con:
        return [json.loads(row['body']) for row in con.execute('SELECT body FROM records WHERE kind=? ORDER BY id', (kind,))]


def merge_companies(con, incoming, clients=False):
    membership = {json.loads(r['body'])['inn'] for r in con.execute("SELECT body FROM records WHERE kind='clients'")}
    if clients:
        for row in incoming:
            member = dict(id=row['inn'], inn=row['inn'], name=row['name'], source=row.get('clientSource') or row['source'], demo=False)
            put(con, 'clients', member)
            membership.add(row['inn'])
    for row in incoming:
        previous = con.execute("SELECT body FROM records WHERE kind='companies' AND id=?", (row['id'],)).fetchone()
        old = json.loads(previous['body']) if previous else {}
        merged = dict(old, **row)
        for key in ['revenue', 'growth', 'lat', 'lon', 'employees', 'statusSource', 'statusCheckedAt']:
            if row.get(key) is None and old.get(key) is not None: merged[key] = old[key]
        if row.get('lat') is None and old.get('lat') is not None:
            merged['geoAccuracy'] = old.get('geoAccuracy','provided')
        if row.get('activityStatus') == 'unknown' and old.get('activityStatus') in ('active', 'terminated'):
            for key in ['activityStatus', 'terminationDate', 'statusSource', 'statusCheckedAt']:
                merged[key] = old.get(key)
        if old.get('registryAsOf') and (clients or (row.get('registryAsOf') or '') < old['registryAsOf']):
            for key in ['mspStatus','mspIncludedAt','mspExcludedAt','registryAsOf','registrySource']:
                merged[key] = old.get(key)
        if row['inn'] in membership or old.get('client') is True:
            merged.update(client=True, clientMatch='confirmed', clientSource=row.get('clientSource') or old.get('clientSource'))
        elif membership:
            merged.update(client=False, clientMatch='not_in_uploaded_base')
        else:
            merged.update(client=None, clientMatch='unknown')
        put(con, 'companies', geocoding.geocode_record(merged, 'companies'))
    if clients:
        # Match registry records imported earlier without overwriting registry fields.
        for record in con.execute("SELECT id,body FROM records WHERE kind='companies'").fetchall():
            row = json.loads(record['body'])
            if row.get('demo'): continue
            row['client'] = row['inn'] in membership
            row['clientMatch'] = 'confirmed' if row['client'] else 'not_in_uploaded_base'
            put(con, 'companies', row)


def load_facts(raw, filename):
    parsed = registry.parse_facts(raw, filename)
    with connect() as con:
        merge_companies(con, parsed['companies'], clients=True)
        for project in parsed['projects']: put(con, 'projects', geocoding.geocode_record(project, 'projects'))
        put(con, 'pricing', parsed['pricing'])
    return dict(imported=len(parsed['companies']), companies=len(parsed['companies']), projects=len(parsed['projects']), equipment=len(parsed['pricing']['items']), message='Клиенты, проекты и базовые цены из фактуры сохранены. Отсутствующие координаты, выручка и стадия проектов не выдумываются.')


def filtered(kind, region='', industry='', q='', stage='', need='', client='', minRevenue=0, minGrowth=-100, demo='', bounds='', registryStatus='current', locationMode='registration', crossLocation='', siteConfidence='confirmed'):
    zone = [float(n) for n in bounds.split(',')] if bounds else []
    if zone and (len(zone) != 4 or not all(math.isfinite(n) for n in zone)):
        raise ValueError('bounds')
    result = []
    site_index = {}
    if kind == 'companies':
        for site in records('sites'): site_index.setdefault(site['inn'], []).append(site)
    for r in records(kind):
        if kind == 'companies':
            r = registry.enrich_company(r, site_index, locationMode, region, siteConfidence)
            if not r.pop('_regionMatches'): continue
            if registryStatus == 'current' and (r.get('mspStatus') == 'excluded' or r.get('activityStatus') == 'terminated'): continue
            if registryStatus == 'excluded' and r.get('mspStatus') != 'excluded': continue
            if registryStatus == 'terminated' and r.get('activityStatus') != 'terminated': continue
            if crossLocation == 'region' and not r['crossRegion']: continue
            if crossLocation == 'locality' and not r['crossLocality']: continue
            if crossLocation == 'candidates' and not r['siteCandidates']: continue
        elif region and r.get('region') != region: continue
        if industry and r.get('industry') != industry: continue
        if stage and r.get('stage') != stage: continue
        if need and r.get('need') != need: continue
        if client:
            value = r.get('client')
            state = 'unknown' if value is None else str(value).lower()
            if state != client: continue
        if q and q.casefold() not in ' '.join(str(v) for v in r.values()).casefold(): continue
        if kind == 'companies':
            if float(minRevenue) > 0 and (r.get('revenue') is None or r['revenue'] < float(minRevenue)): continue
            if float(minGrowth) > -100 and (r.get('growth') is None or r['growth'] < float(minGrowth)): continue
        if demo and str(r.get('demo', False)).lower() != demo: continue
        if zone:
            locations = r.get('mapLocations', [r])
            if not any(p.get('lat') is not None and p.get('lon') is not None and zone[0] <= p['lat'] <= zone[2] and zone[1] <= p['lon'] <= zone[3] for p in locations): continue
        result.append(r)
    return result


def query_records(kind, request):
    allowed = {'region', 'industry', 'q', 'stage', 'need', 'client', 'minRevenue', 'minGrowth', 'demo', 'bounds', 'registryStatus', 'locationMode', 'crossLocation', 'siteConfidence'}
    if kind == 'companies': allowed -= {'stage', 'need'}
    if kind == 'projects': allowed -= {'client', 'minRevenue', 'minGrowth'}
    args = {k: v for k, v in request.query_params.items() if k in allowed}
    try:
        return filtered(kind, **args)
    except ValueError:
        raise HTTPException(422, 'Некорректное числовое значение фильтра')


initialize()


@app.get('/health')
def health():
    return {'status': 'ok', 'service': 'analytics'}


@app.get('/catalog')
def catalog():
    return dict(regions=list(REGIONS), industries=INDUSTRIES, stages=STAGES, needs=TYPES)


@app.get('/companies')
def companies(request: Request, limit: int = Query(default=1000, ge=1, le=5000), offset: int = Query(default=0, ge=0)):
    return query_records('companies', request)[offset:offset+limit]


@app.get('/projects')
def projects(request: Request):
    return sorted(query_records('projects', request), key=lambda p: p['date'])


@app.get('/insights')
def insights(request: Request):
    rows = query_records('companies', request)
    groups = {}
    for r in rows:
        key = (r['region'], r['industry'])
        group = groups.setdefault(key, dict(region=key[0], industry=key[1], companies=0, clients=0, revenue=0, growth=0))
        group['companies'] += 1
        group['clients'] += int(r.get('client') is True)
        group['potentialClients'] = group.get('potentialClients', 0) + int(r.get('client') is False)
        group['revenue'] += r.get('revenue') or 0
        group['growth'] += r.get('growth') or 0
        group['revenueCoverage'] = group.get('revenueCoverage', 0) + int(r.get('revenue') is not None)
        group['growthCoverage'] = group.get('growthCoverage', 0) + int(r.get('growth') is not None)
    gaps = []
    for g in groups.values():
        g['growth'] = round(g['growth']/g['growthCoverage'], 1) if g['growthCoverage'] else None
        g['penetration'] = round(g['clients']/g['companies']*100)
        g['score'] = round((g['potentialClients']/g['companies'])*max(0, (g['growth'] or 0)+10)*math.log1p(g['revenue']), 1) if g['revenueCoverage'] and g['growthCoverage'] else g['potentialClients']
        g['basis'] = 'financials' if g['revenueCoverage'] and g['growthCoverage'] else 'company_count'
        if not g['revenueCoverage']: g['revenue'] = None
        if g['potentialClients']: gaps.append(g)
    gaps.sort(key=lambda g: (g['basis'] == 'financials', g['score']), reverse=True)
    revenue_values = [r['revenue'] for r in rows if r.get('revenue') is not None]
    growth_values = [r['growth'] for r in rows if r.get('growth') is not None]
    return dict(companies=len(rows), clients=sum(r.get('client') is True for r in rows), unknownClients=sum(r.get('client') is None for r in rows), revenue=round(sum(revenue_values),1) if revenue_values else None, growth=round(sum(growth_values)/len(growth_values),1) if growth_values else None, revenueCoverage=len(revenue_values), growthCoverage=len(growth_values), registryCoverage=sum(bool(r.get('registryAsOf')) for r in rows), locatedCompanies=sum(bool(r.get('mapLocations')) for r in rows), crossRegion=sum(r.get('crossRegion',False) for r in rows), gaps=gaps[:8], demoCount=sum(r['demo'] for r in rows))


class Estimate(BaseModel):
    type: Literal['ИТ-инфраструктура', 'Видеонаблюдение', 'Кибербезопасность', 'АСУ ТП', 'Связь']
    region: str
    units: int = Field(ge=1, le=100000)
    months: int = Field(ge=1, le=120)
    financing: Literal['Единовременно', 'Рассрочка', 'Лизинг'] = 'Единовременно'
    equipment: float = Field(ge=0, le=1e12)
    software: float = Field(ge=0, le=1e12)
    labor: float = Field(ge=0, le=1e12)
    urgent: bool = False
    unique: bool = False
    customer: str = Field(default='', max_length=200)


def estimate(data: Estimate):
    if data.region not in REGIONS:
        raise HTTPException(422, 'Неизвестный регион')
    analogs = [dict(d, distance=abs(math.log(d['units']/data.units)) + abs(d['months']-data.months)/12 + (0 if d['region'] == data.region else .15)) for d in records('deals') if d['type'] == data.type]
    # Never mix demonstration prices into estimates based on imported historical deals.
    real = [d for d in analogs if not d['demo']]
    analogs = sorted(real or analogs, key=lambda d: d['distance'])[:5]
    if not analogs:
        raise HTTPException(422, 'Для выбранного типа нет аналогов. Импортируйте историю сделок.')
    weight = sum(1/(1+d['distance']) for d in analogs)
    unit_price = sum(d['price']/d['units']/(1+d['distance']) for d in analogs)/weight
    regional = REGIONS[data.region][2] 
    if records('pricing') and data.region in ('Ростовская область', 'Краснодарский край', 'Волгоградская область'):
        regional = max(regional, 1.1)
    urgency = 1.15 if data.urgent or data.months <= 2 else 1
    unique = 1.2 if data.unique else 1
    financing = {'Единовременно': 1, 'Рассрочка': 1.08, 'Лизинг': 1.12}[data.financing]
    cost = (data.equipment+data.software+data.labor*regional)*urgency*unique
    if cost <= 0:
        raise HTTPException(422, 'Укажите ненулевую себестоимость')
    market = unit_price*data.units*regional*urgency*unique*financing
    minimum = max(cost/.9*financing, market*.85)
    optimum = max(cost/.75*financing, market)
    maximum = max(optimum*1.18, minimum*1.2)
    return dict(minimum=round(minimum), optimum=round(optimum), maximum=round(maximum), cost=round(cost*financing), margin=round((optimum-cost*financing)/optimum*100, 1), analogs=analogs, regional=regional, urgency=urgency, unique=unique, financing=financing, demo=all(d['demo'] for d in analogs), confidence='Средняя' if len(analogs)>=5 and analogs[0]['distance']<.6 else 'Низкая', method='Средневзвешенная цена единицы 5 ближайших аналогов; нижняя граница не ниже себестоимости / 0,9; оптимум не ниже себестоимости / 0,75. Цены без НДС.', input=data.model_dump())


@app.post('/estimate')
def calculate(data: Estimate):
    return estimate(data)


def csv_response(rows, filename):
    stream = io.StringIO(newline='')
    if rows:
        columns = list(dict.fromkeys(k for r in rows for k in r))
        writer = csv.DictWriter(stream, fieldnames=columns, delimiter=';')
        writer.writeheader()
        for row in rows:
            # Avoid spreadsheet formula execution on CRM export.
            writer.writerow({k: ("'"+v if isinstance(v, str) and v.startswith(('=', '+', '-', '@', '\t', '\r')) else v) for k, v in row.items()})
    return Response(('\ufeff'+stream.getvalue()).encode('utf-8'), media_type='text/csv; charset=utf-8', headers={'Content-Disposition': f'attachment; filename="{filename}.csv"'})


@app.get('/export/{kind}')
def export(kind: Literal['companies', 'projects', 'deals'], request: Request, crm: bool = False):
    rows = query_records(kind, request)
    if crm and kind == 'companies':
        rows = [dict(
            COMPANY_TITLE=r.get('name'), COMPANY_INN=r.get('inn'), ADDRESS=r.get('address'), REGION=r.get('region'), INDUSTRY=r.get('industry'), REVENUE_MLN_RUB=r.get('revenue'), IS_CLIENT=r.get('client'), SOURCE=r.get('source')) for r in rows]
    return csv_response(rows, f'{kind}-crm' if crm else kind)


@app.post('/estimate/export/{format}')
def export_estimate(format: Literal['csv', 'docx', 'doc'], data: Estimate):
    result = estimate(data)
    if format == 'csv':
        return csv_response([dict(data.model_dump(), minimum=result['minimum'], optimum=result['optimum'], maximum=result['maximum'], margin=result['margin'], demo=result['demo'])], 'estimate')
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading('Коммерческое предложение', 0)
    doc.add_paragraph(f'Ростелеком · {date.today().strftime("%d.%m.%Y")}')
    if data.customer: doc.add_paragraph(f'Заказчик: {data.customer}')
    doc.add_paragraph(f'{data.type}. Регион: {data.region}. Объём: {data.units} ед. Срок: {data.months} мес. Финансирование: {data.financing}.')
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Light Shading Accent 1'
    for label, value in [('Минимальная цена', result['minimum']), ('Рекомендуемая цена', result['optimum']), ('Верхняя граница', result['maximum'])]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = f'{value:,.0f} руб. без НДС'.replace(',', ' ')
    doc.add_paragraph('Предварительная оценка. Финальная стоимость уточняется после обследования объекта. Срок действия предложения: 30 дней.')
    if result['demo']: doc.add_paragraph('Расчёт на демонстрационной истории сделок. Не использовать как рыночное обоснование цены.')
    doc.add_heading('Основание расчёта', 1)
    doc.add_paragraph(result['method'])
    for d in result['analogs']: doc.add_paragraph(f'{d["name"]}: {d["price"]:,.0f} руб.; {d["units"]} ед.; {d["source"]}.')
    if format == 'doc':
        import html
        paragraphs = ''.join(f'<p>{html.escape(p.text)}</p>' for p in doc.paragraphs)
        table_html = '<table border="1" cellpadding="4">' + ''.join('<tr>' + ''.join(f'<td>{html.escape(c.text)}</td>' for c in row.cells) + '</tr>' for row in doc.tables[0].rows) + '</table>'
        content = f'<html><head><meta charset="utf-8"></head><body>{paragraphs}{table_html}</body></html>'
        return Response(content.encode('utf-8'), media_type='application/msword',
                        headers={'Content-Disposition': 'attachment; filename="proposal.doc"'})
    stream = io.BytesIO()
    doc.save(stream)
    return Response(stream.getvalue(), media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                    headers={'Content-Disposition': 'attachment; filename="proposal.docx"'})


@app.get('/data-status')
def data_status():
    return [dict(kind=kind, count=len(rows := records(kind)), demo=sum(r.get('demo', False) for r in rows), sources=sorted(set(r.get('source','') for r in rows))) for kind in ['companies', 'projects', 'deals', 'clients', 'sites']]


@app.get('/pricing-catalog')
def pricing_catalog():
    rows = records('pricing')
    return rows[0] if rows else None


@app.get('/registry-status')
def registry_status():
    rows = [r for r in records('companies') if not r.get('demo')]
    return dict(companies=len(rows), matchedClients=sum(r.get('client') is True for r in rows), unmatched=sum(r.get('client') is False for r in rows), unknownClients=sum(r.get('client') is None for r in rows), excludedMsp=sum(r.get('mspStatus')=='excluded' for r in rows), terminated=sum(r.get('activityStatus')=='terminated' for r in rows), sites=len(records('sites')), sourceUrl=registry.FNS_URL)


@app.get('/template/sites')
def sites_template():
    return csv_response([dict(inn='',region='Ростовская область',locality='г. Азов',address='',source='',sourceUrl='',observedAt='',status='candidate',lat='',lon='',geoAccuracy='provided')], 'operating-sites-template')


def _source_import(kind, raw, filename, asOf):
    try:
        if kind == 'facts': return load_facts(raw, Path(filename).name)
        if kind == 'sites':
            rows = registry.parse_sites(raw, Path(filename).name)
            with connect() as con:
                known = {json.loads(r['body'])['inn'] for r in con.execute("SELECT body FROM records WHERE kind='companies'")}
                if any(r['inn'] not in known for r in rows): raise ValueError('для площадки не найден ИНН в базе компаний; сначала импортируйте компанию')
                for row in rows: put(con, 'sites', row)
        else:
            rows = registry.parse_companies(raw, Path(filename).name, is_client=kind=='clients')
            if kind == 'fns':
                snapshot = registry.iso_date(asOf)
                if not snapshot or snapshot > date.today().isoformat(): raise ValueError('укажите действительную дату среза ФНС asOf в формате YYYY-MM-DD')
                for row in rows: row.update(registryAsOf=snapshot, registrySource='ФНС · Реестр МСП')
            with connect() as con: merge_companies(con, rows, clients=kind=='clients')
        return dict(imported=len(rows), excludedMsp=sum(r.get('mspStatus')=='excluded' for r in rows), message='Импорт завершён. ИНН сопоставлены точно; исключение из МСП не считается прекращением деятельности.')
    except (ValueError, TypeError) as exc:
        raise HTTPException(422, str(exc)+'. Импорт отменён целиком.')


@app.post('/source-import/{kind}')
async def source_import(kind: Literal['clients','fns','sites','facts'], request: Request, filename: str = 'upload.csv', asOf: str = ''):
    raw = await request.body()
    if len(raw) > 25_000_000: raise HTTPException(413, 'Максимальный размер исходного файла — 25 МБ; разбейте большую выгрузку на части')
    return await run_in_threadpool(_source_import, kind, raw, filename, asOf)


@app.get('/template/{kind}')
def template(kind: Literal['companies', 'projects', 'deals'], format: Literal['csv', 'docx'] = 'csv'):
    rows = records(kind)
    if not rows:
        raise HTTPException(404, 'Нет записей, по которым можно построить шаблон')
    sample = rows[0].copy()
    for key in ['id', 'demo', 'updatedAt', 'distance']: sample.pop(key, None)
    sample['source'] = 'Название реального источника'
    sample['sourceUrl'] = 'https://example.org/record'
    if format == 'docx':
        from docx import Document
        from docx.shared import Cm, Pt
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Cm(42), Cm(29.7)
        section.left_margin = section.right_margin = Cm(1)
        document.styles['Normal'].font.size = Pt(9)
        document.add_heading('Шаблон импорта данных', 0)
        document.add_paragraph('Первая строка каждой таблицы — названия полей. Заполните строки по примеру. '
                               'Названия полей должны совпадать с шаблоном; сохраняйте ИНН как текст. '
                               'Не объединяйте ячейки. Для удобства поля примера разделены на строки: '
                               'перенесите нужные значения в таблицу данных ниже.')
        # Keep the reference outside the data table so only actual records get imported.
        reference = document.add_paragraph()
        for key, value in sample.items():
            reference.add_run(f'{key}: {value}\n')
        document.add_paragraph('Таблица данных (вставьте записи, затем удалите эту инструкцию):')
        table = document.add_table(rows=1, cols=len(sample))
        table.style = 'Table Grid'
        for cell, key in zip(table.rows[0].cells, sample): cell.text = key
        document.add_paragraph('')  # Preserve a paragraph after the table when saving as legacy DOC.
        stream = io.BytesIO()
        document.save(stream)
        return Response(stream.getvalue(), media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                        headers={'Content-Disposition': f'attachment; filename="{kind}-template.docx"'})
    return csv_response([sample], f'{kind}-template')


@app.post('/import/{kind}')
async def import_csv(kind: Literal['companies', 'projects', 'deals'], request: Request, filename: str = 'upload.csv'):
    raw = await request.body()
    if len(raw) > 5_000_000: raise HTTPException(413, 'Максимальный размер файла — 5 МБ')
    extension = Path(filename).suffix.lower()
    if extension in ('.doc', '.docx'):
        try:
            raw = await run_in_threadpool(word_import.word_to_csv, raw, filename)
        except ValueError as exc:
            raise HTTPException(422, str(exc) + '. Ни одна запись не импортирована.') from exc
    elif extension != '.csv':
        raise HTTPException(422, 'Выберите CSV, DOC или DOCX')
    try:
        text = raw.decode('utf-8-sig')
        delimiter = ';' if ';' in text.splitlines()[0] else ','
        rows = list(csv.DictReader(io.StringIO(text), delimiter=delimiter))
    except (UnicodeError, IndexError, csv.Error):
        raise HTTPException(422, 'Не удалось прочитать CSV. Используйте UTF-8 и шаблон.')
    if not 1 <= len(rows) <= 10000: raise HTTPException(422, 'Допустимо от 1 до 10000 строк')
    clean, seen = [], set()
    import hashlib
    for index, row in enumerate(rows, 2):
        try:
            def required(key):
                value = str(row.get(key, '')).strip()
                if not value or len(value) > 1000: raise ValueError(f'поле {key} обязательно, до 1000 символов')
                return value
            def number(key, low, high):
                value = float(required(key).replace(',', '.'))
                if not math.isfinite(value) or not low <= value <= high: raise ValueError(f'{key}: диапазон {low}…{high}')
                return value
            r = dict(name=required('name'), region=required('region'), source=required('source'), sourceUrl=str(row.get('sourceUrl', '')).strip(), updatedAt=date.today().isoformat(), demo=False)
            if r['region'] not in REGIONS: raise ValueError('регион отсутствует в справочнике')
            if r['sourceUrl'] and not r['sourceUrl'].startswith(('https://', 'http://')): raise ValueError('sourceUrl должен быть HTTP(S) URL')
            if kind == 'companies':
                inn = required('inn')
                if not inn.isdigit() or len(inn) not in (10,12): raise ValueError('ИНН: 10 или 12 цифр')
                client = required('client').lower()
                if client not in ('true', 'false', '1', '0', 'да', 'нет'): raise ValueError('client: допустимы true/false, 1/0, да/нет')
                r.update(inn=inn, industry=required('industry'), revenue=number('revenue',0,1e12), growth=number('growth',-100,1e6), client=client in ('true','1','да'), address=required('address'), employees=int(number('employees',0,1e8)))
                identity = inn
            elif kind == 'projects':
                r.update(initiator=required('initiator'), industry=required('industry'), stage=required('stage'), need=required('need'), investment=number('investment',0,1e12), date=required('date'))
                date.fromisoformat(r['date'])
                if r['stage'] not in STAGES or r['need'] not in TYPES: raise ValueError('неизвестная стадия или потребность')
                identity = r['name']+r['region']+r['initiator']
            else:
                r.update(type=required('type'), units=int(number('units',1,100000)), months=int(number('months',1,120)), price=number('price',1,1e15), cost=number('cost',0,1e15), date=required('date'))
                date.fromisoformat(r['date'])
                if r['type'] not in TYPES: raise ValueError('неизвестный тип сделки')
                identity = r['name']+r['region']+r['date']
            if kind != 'deals':
                if r['industry'] not in INDUSTRIES: raise ValueError('неизвестная отрасль')
                r.update(lat=number('lat',-90,90), lon=number('lon',-180,180))
            r['id'] = 'real-'+hashlib.sha256(identity.encode()).hexdigest()[:20]
            if r['id'] in seen: raise ValueError('дубликат в файле')
            seen.add(r['id'])
            clean.append(r)
        except (ValueError, TypeError, AttributeError) as exc:
            raise HTTPException(422, f'Строка {index}: {exc}. Ни одна строка не импортирована.')
    with connect() as con:
        for row in clean: put(con, kind, row)
    return dict(imported=len(clean), message='Данные проверены и сохранены; повторный импорт обновляет записи.')
